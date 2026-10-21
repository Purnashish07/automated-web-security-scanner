from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Page setup
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

# Default font
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)

# Cover page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\n\nAUTOMATED WEB SECURITY SCANNER")
r.bold = True
r.font.size = Pt(22)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Python-Based Website Security Configuration Audit")
r.italic = True
r.font.size = Pt(13)

doc.add_paragraph("\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Project Report\n").bold = True
p.add_run("IBM SkillsBuild Cyber Security BootCamp 2026\n\n")
p.add_run("Submitted by: Purnashish Atroliya\n")
p.add_run("Madhav Institute of Technology and Science, Gwalior\n")
p.add_run("BharatCares")

doc.add_page_break()

doc.add_heading("1. Abstract", level=1)
doc.add_paragraph(
    "The Automated Web Security Scanner is an educational Python-based tool "
    "that checks selected website security configurations. It evaluates HTTPS "
    "availability, common HTTP security headers, cookie attributes, and basic "
    "Cross-Origin Resource Sharing (CORS) settings. A Flask web interface allows "
    "a user to submit an authorized website URL and review the findings. The "
    "application can export a JSON report for documentation and further review. "
    "The tool identifies configuration indicators and does not prove that a "
    "website is vulnerable or secure."
)

doc.add_heading("2. Problem Statement", level=1)
doc.add_paragraph(
    "Website security settings can be difficult for beginners to inspect "
    "manually. Missing or weak HTTP security headers and cookie attributes may "
    "increase exposure to certain web risks. This project provides a simple, "
    "readable first-pass review of selected security configurations."
)

doc.add_heading("3. Objectives", level=1)
for item in [
    "Check whether a target URL uses HTTPS.",
    "Inspect common HTTP security headers.",
    "Review cookie security attributes when cookies are present.",
    "Identify basic CORS configuration indicators.",
    "Display findings in a browser-based interface.",
    "Export scan results as a JSON report.",
    "Promote responsible testing of authorized targets only.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("4. Technology Stack", level=1)
for item in [
    "Python: scanner logic and request handling.",
    "Flask: local web application and scan API.",
    "Requests: HTTP/HTTPS requests.",
    "HTML, CSS and JavaScript: browser interface.",
    "python-docx: generation of this project report.",
    "JSON: structured scan report export.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("5. System Design and Workflow", level=1)
doc.add_paragraph(
    "User enters a URL → Flask validates the input → scanner requests the target "
    "and inspects the received response → checks are classified as Pass, Warning "
    "or Review → results are displayed in the interface → user can export JSON."
)
doc.add_paragraph(
    "The scanner includes checks for HTTPS, Strict-Transport-Security, "
    "Content-Security-Policy, X-Content-Type-Options, X-Frame-Options, "
    "Referrer-Policy, cookie attributes and basic CORS response headers."
)

doc.add_heading("6. Implementation", level=1)
doc.add_paragraph(
    "The application uses a Flask backend and a browser-based frontend. "
    "The scanner sends HTTP requests and evaluates selected response properties. "
    "The interface presents the results and provides JSON export. Input validation "
    "is used to restrict scan targets to public destinations. Testing should be "
    "performed only on websites for which permission has been granted."
)

doc.add_heading("7. Testing and Observed Results", level=1)
doc.add_paragraph(
    "During a sample run, the application returned eight checks: two Pass, "
    "four Warning and two Review results. HTTPS and HSTS were marked Pass. "
    "CSP, X-Content-Type-Options, X-Frame-Options and Referrer-Policy were "
    "reported as Warning. CORS was marked Review because a wildcard origin "
    "was observed, and cookies were marked Review because no Set-Cookie header "
    "was observed in that response. These results describe the tested response "
    "only; they are not a comprehensive vulnerability assessment."
)
doc.add_paragraph(
    "Note: Re-run the scanner on an authorized target before submission and "
    "replace these sample observations if your final test produces different results."
)

doc.add_heading("8. Limitations", level=1)
for item in [
    "Checks only a limited set of HTTP security configurations.",
    "Does not perform penetration testing or exploit vulnerabilities.",
    "A missing header does not automatically mean the website is exploitable.",
    "Cookie findings depend on whether the response sets cookies.",
    "CORS findings require endpoint and application-context review.",
    "Redirect and DNS-based protections are not a substitute for a hardened production scanner.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("9. Future Enhancements", level=1)
for item in [
    "Add unit tests and automated regression tests.",
    "Improve redirect validation and network-target protections.",
    "Add timestamped reports and clearer severity explanations.",
    "Provide downloadable HTML and PDF reports.",
    "Add configurable checks and a test suite using a local test server.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("10. Ethical and Responsible Use", level=1)
doc.add_paragraph(
    "Use this tool only on websites you own or have explicit permission to assess. "
    "Do not use it to bypass access controls, overload services, or access data "
    "without authorization. Findings should be verified by a qualified reviewer "
    "before any security conclusion is made."
)

doc.add_heading("11. Conclusion", level=1)
doc.add_paragraph(
    "This project demonstrates a beginner-accessible approach to reviewing "
    "selected website security response settings using Python and Flask. It "
    "provides an interactive interface and machine-readable JSON output while "
    "emphasizing authorization, limitations and responsible interpretation."
)

doc.add_heading("12. References", level=1)
for item in [
    "Python documentation: https://docs.python.org/3/",
    "Flask documentation: https://flask.palletsprojects.com/",
    "Requests documentation: https://requests.readthedocs.io/",
    "OWASP Secure Headers Project: https://owasp.org/www-project-secure-headers/",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.save("Project_Report.docx")
print("Created Project_Report.docx")
